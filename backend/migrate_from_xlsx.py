"""
구글 시트(엑셀로 내려받은 파일)의 데이터를 Supabase로 옮깁니다.

  python -I -X utf8 backend/migrate_from_xlsx.py --src <엑셀 폴더>                 # 전체 게임
  python -I -X utf8 backend/migrate_from_xlsx.py --src <엑셀 폴더> --appid 4126220  # 한 게임만 (시험용)

- 여러 번 실행해도 안전합니다. 같은 리뷰는 다시 들어가지 않습니다(중복은 건너뜀).
- 원본 엑셀은 읽기만 하고 절대 바꾸지 않습니다.
- 플레이타임은 원본에 값이 없었기 때문에(전부 0) 비워 두고, 새 수집기가 나중에 채웁니다.
"""
import argparse, glob, json, os, re, sys, time
from datetime import date, datetime, timezone

import truststore; truststore.inject_into_ssl()
import openpyxl, psycopg, requests
from dotenv import load_dotenv

here = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(here, "..", ".env"))
UTC = timezone.utc


# ───────────── 값 변환 도우미 ─────────────
def s(v):
    """셀 값을 문자열로. 비어 있으면 None. DB가 싫어하는 널 문자는 제거."""
    if v is None:
        return None
    t = str(v).replace("\x00", "").strip()
    return t if t != "" else None

def to_bool(v):
    if isinstance(v, bool):
        return v
    t = s(v)
    return None if t is None else t.lower() in ("true", "1", "yes")

def to_int(v, default=None):
    t = s(v)
    if t is None:
        return default
    try:
        return int(float(t))
    except ValueError:
        return default

def to_num(v):
    t = s(v)
    if t is None:
        return None
    try:
        return float(t)
    except ValueError:
        return None

def to_dt(v):
    """'2026-04-19T11:28:57Z' 또는 '2026-03-31 03:00:00' (모두 UTC로 간주)"""
    if isinstance(v, datetime):
        return v.replace(tzinfo=v.tzinfo or UTC)
    t = s(v)
    if not t:
        return None
    try:
        d = datetime.fromisoformat(t.replace("Z", "+00:00"))
        return d if d.tzinfo else d.replace(tzinfo=UTC)
    except ValueError:
        return None

def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = s(v)
    if not t:
        return None
    try:
        return date.fromisoformat(t[:10])
    except ValueError:
        return None

def parse_release(text):
    t = s(text)
    if not t:
        return None
    for fmt in ("%Y-%m-%d", "%b %d, %Y", "%d %b, %Y", "%d %b %Y", "%b %d %Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    m = re.search(r"(\d{4})년\s*(\d{1,2})월\s*(\d{1,2})일", t)
    return date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None

def jload(v, default):
    t = s(v)
    if not t:
        return default
    try:
        return json.loads(t)
    except ValueError:
        return default

def rows_of(ws):
    """헤더를 이름표로 쓰는 딕셔너리 줄 단위 읽기"""
    it = ws.iter_rows(values_only=True)
    hdr = next(it, None)
    if not hdr:
        return
    names = [str(h).strip() if h is not None else "" for h in hdr]
    for r in it:
        if r is None or all(c is None for c in r):
            continue
        yield dict(zip(names, r))


# ───────────── 게임 정보 ─────────────
def steam_header_image(appid):
    try:
        r = requests.get("https://store.steampowered.com/api/appdetails",
                         params={"appids": appid, "filters": "basic"}, timeout=20)
        d = r.json().get(str(appid), {})
        return d.get("data", {}).get("header_image") if d.get("success") else None
    except Exception:
        return None

def import_games(conn, master_path, only):
    wb = openpyxl.load_workbook(master_path, read_only=True)
    n = 0
    appids = []
    for g in rows_of(wb["games"]):
        appid = to_int(g.get("appid"))
        if not appid or (only and appid not in only):
            continue
        status = {"error_pool_empty": "error"}.get(s(g.get("status")) or "", s(g.get("status")) or "active")
        if status not in ("collecting", "active", "archived", "error"):
            status = "active"
        thumb = steam_header_image(appid) or s(g.get("thumbnail"))
        time.sleep(1.2)
        rel_text = s(g.get("release_date"))
        genres = [x.strip() for x in (s(g.get("genres")) or "").split(",") if x.strip()]
        langs = [x.strip() for x in (s(g.get("top_languages")) or "").split(",") if x.strip()]
        conn.execute("""
            insert into games (appid, name, thumbnail, status, release_date, release_date_text, genres, developer,
                               publisher, price, is_free, is_early_access, steam_total_reviews, steam_positive_rate,
                               top_languages, ai_enabled, collection_started_at)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,coalesce(%s, now()))
            on conflict (appid) do update set
              name = excluded.name, thumbnail = excluded.thumbnail, status = excluded.status,
              release_date = excluded.release_date, release_date_text = excluded.release_date_text,
              genres = excluded.genres, developer = excluded.developer, publisher = excluded.publisher,
              price = excluded.price, is_free = excluded.is_free, is_early_access = excluded.is_early_access,
              steam_total_reviews = excluded.steam_total_reviews, steam_positive_rate = excluded.steam_positive_rate,
              top_languages = excluded.top_languages, ai_enabled = excluded.ai_enabled
        """, (appid, s(g.get("name")) or str(appid), thumb, status, parse_release(rel_text), rel_text, genres,
              s(g.get("developer")), s(g.get("publisher")), s(g.get("price")), to_bool(g.get("is_free")),
              to_bool(g.get("is_early_access")), to_int(g.get("total_reviews_count")),
              to_num(g.get("steam_positive_rate")), langs, bool(to_bool(g.get("ai_approved"))) or bool(s(g.get("ai_briefing"))),
              to_dt(g.get("collection_started_at"))))
        for kind, col in (("briefing", "ai_briefing"), ("trend", "sentiment_trend_comment"),
                          ("ccu", "ccu_peaktime_comment"), ("language_cross", "language_cross_comment"),
                          ("lifecycle", "lifecycle_comment")):
            body = s(g.get(col))
            if body:
                conn.execute("""insert into game_reports (appid, kind, body) values (%s,%s,%s)
                                on conflict (appid, kind) do update set body = excluded.body""", (appid, kind, body))
        appids.append(appid)
        n += 1
    conn.commit()
    return appids


# ───────────── 리뷰 ─────────────
def import_reviews(conn, appid, files):
    """엑셀 탭마다 읽어 임시 표로 COPY한 뒤, 중복을 건너뛰고 reviews에 넣습니다."""
    src_rows, src_ids, inserted = 0, set(), 0
    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True)
        for ws in wb.worksheets:
            if not ws.title.startswith("reviews_"):
                continue
            conn.execute("""create temp table if not exists _rv (
                recommendationid bigint, voted_up boolean, language text, review text,
                created_at timestamptz, votes_up integer) on commit preserve rows""")
            n = 0
            buf = []

            def flush():
                """모아 둔 줄을 2만 건씩 끊어서 넣습니다 (한 번에 너무 크면 접속이 끊길 수 있음)"""
                nonlocal inserted
                if not buf:
                    return
                conn.execute("truncate _rv")
                with conn.cursor().copy("copy _rv from stdin") as cp:
                    for row in buf:
                        cp.write_row(row)
                cur = conn.execute("""insert into reviews (appid, recommendationid, voted_up, language, review, created_at, votes_up)
                    select %s, recommendationid, voted_up, language, review, created_at, votes_up from _rv
                    on conflict (appid, recommendationid) do nothing""", (appid,))
                inserted += cur.rowcount
                conn.commit()
                buf.clear()

            for r in rows_of(ws):
                rid = to_int(r.get("recommendationid"))
                ts = to_num(r.get("timestamp_created"))
                up = to_bool(r.get("voted_up"))
                if not rid or ts is None or up is None:
                    continue
                src_rows += 1; n += 1; src_ids.add(rid)
                buf.append((rid, up, s(r.get("language")), s(r.get("review")),
                            datetime.fromtimestamp(ts, tz=UTC), to_int(r.get("votes_up"), 0)))
                if len(buf) >= 20000:
                    flush()
            flush()
            print("    %s / %s: 읽음 %d" % (os.path.basename(f)[:40], ws.title, n), flush=True)
        wb.close()
    return src_rows, len(src_ids), inserted


# ───────────── 동접 · 긍정률 · 타임라인 ─────────────
def import_side_tabs(conn, appid, files):
    ccu = rate = ev = an = 0
    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True)
        names = wb.sheetnames
        if "ccu" in names:
            batch = []
            for r in rows_of(wb["ccu"]):
                t = to_dt(r.get("timestamp")); v = to_int(r.get("ccu_value"))
                if t is None or v is None:
                    continue
                batch.append((appid, t, v, "steamdb" if to_bool(r.get("is_archived_gap")) else "live"))
            with conn.cursor() as cur:
                cur.executemany("""insert into ccu_samples (appid, sampled_at, ccu, source) values (%s,%s,%s,%s)
                                   on conflict do nothing""", batch)
            ccu += len(batch)
        if "rate_history" in names:
            for r in rows_of(wb["rate_history"]):
                d = to_date(r.get("date"))
                if d is None:
                    continue
                conn.execute("""insert into rate_history (appid, day, positive_rate, total_reviews) values (%s,%s,%s,%s)
                                on conflict do nothing""", (appid, d, to_num(r.get("positive_rate")), to_int(r.get("total_reviews"))))
                rate += 1
        if "timeline" in names:
            for r in rows_of(wb["timeline"]):
                e, a = import_timeline_row(conn, appid, r)
                ev += e; an += a
        conn.commit()
        wb.close()
    return ccu, rate, ev, an

def import_timeline_row(conn, appid, r):
    et = s(r.get("event_type")); scope = s(r.get("language_scope")) or "all"
    d = to_date(r.get("date")); eid = s(r.get("event_id"))
    if not et or not d or not eid:
        return 0, 0
    if et in ("official", "news", "manual", "free_weekend", "launch"):
        if scope != "all":
            return 0, 0
        conn.execute("""insert into events (appid, event_key, type, event_date, title, title_kr, url, content, ai_summary, sale_text)
                        values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) on conflict (appid, event_key) do nothing""",
                     (appid, eid, et, d, s(r.get("title")) or eid, s(r.get("title_kr")), s(r.get("url")),
                      s(r.get("content")), s(r.get("ai_patch_summary")), s(r.get("sale_text"))))
        return 1, 0
    rate = s(r.get("sentiment_rate")); sparse = rate == "sparse"
    common = dict(title=s(r.get("title_kr")) or s(r.get("title")), rate=None if sparse else to_num(rate),
                  cnt=to_int(r.get("review_count")), sparse=sparse, summ=s(r.get("ai_reaction_summary")),
                  kw=json.dumps(jload(r.get("top_keywords"), []), ensure_ascii=False),
                  tr=json.dumps(jload(r.get("top_reviews"), []), ensure_ascii=False))
    if et == "monthly_summary":
        kind, pkey, end = "monthly", d.strftime("%Y-%m"), None
        extra = {"patch_summary": s(r.get("ai_patch_summary")) or ""}
    elif et == "weekly_summary":
        kind, pkey, end = "weekly", d.isoformat(), to_date(r.get("date_end"))
        extra = {}
    elif et == "sentiment_shift":
        kind, pkey, end = "shift", d.isoformat(), to_date(r.get("date_end"))
        extra = {"sentiment_before": to_num(r.get("sentiment_before")), "sentiment_delta": to_num(r.get("sentiment_delta")),
                 "direction": s(r.get("direction")), "confidence": s(r.get("confidence")),
                 "linked_event_ids": jload(r.get("linked_event_ids"), []), "is_official_confirmed": to_bool(r.get("is_official_confirmed"))}
    else:
        return 0, 0
    conn.execute("""insert into analyses (appid, kind, period_key, scope, period_start, period_end, title, sentiment_rate,
                      review_count, is_sparse, reaction_summary, keywords, top_reviews, extra)
                    values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::jsonb,%s::jsonb,%s::jsonb)
                    on conflict (appid, kind, period_key, scope) do update set
                      sentiment_rate = excluded.sentiment_rate, review_count = excluded.review_count,
                      is_sparse = excluded.is_sparse, reaction_summary = excluded.reaction_summary,
                      keywords = excluded.keywords, top_reviews = excluded.top_reviews, extra = excluded.extra""",
                 (appid, kind, pkey, scope, d, end, common["title"], common["rate"], common["cnt"], common["sparse"],
                  common["summ"], common["kw"], common["tr"], json.dumps(extra, ensure_ascii=False)))
    return 0, 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--appid", type=int, nargs="*")
    a = ap.parse_args()
    only = set(a.appid or [])
    master = os.path.join(a.src, "Steam-Pickaxe-Data.xlsx")
    t0 = time.time()
    with psycopg.connect(os.environ["SUPABASE_DB_URL"], connect_timeout=30, autocommit=False,
                         keepalives=1, keepalives_idle=30, keepalives_interval=10, keepalives_count=5) as conn:
        appids = import_games(conn, master, only)
        print("게임 %d개 등록/갱신" % len(appids), flush=True)
        report = []
        for appid in appids:
            files = sorted(glob.glob(os.path.join(a.src, "RAW_REVIEWS_%d_*.xlsx" % appid)))
            name = conn.execute("select name from games where appid=%s", (appid,)).fetchone()[0]
            print("\n[%d] %s — 파일 %d개" % (appid, name, len(files)), flush=True)
            if not files:
                report.append((appid, name, 0, 0, 0, 0)); continue
            rows, uniq, ins = import_reviews(conn, appid, files)
            ccu, rate, ev, an = import_side_tabs(conn, appid, files)
            conn.execute("select refresh_review_daily(%s)", (appid,))
            conn.execute("select refresh_game_counts(%s)", (appid,))
            conn.commit()
            in_db = conn.execute("select count(*) from reviews where appid=%s", (appid,)).fetchone()[0]
            print("  리뷰: 원본 줄수 %d, 고유 %d, DB %d (이번에 새로 넣음 %d) | 동접 %d, 긍정률 %d, 이벤트 %d, 분석 %d"
                  % (rows, uniq, in_db, ins, ccu, rate, ev, an), flush=True)
            report.append((appid, name, rows, uniq, in_db, ins))
    print("\n──── 대조 결과 ────")
    ok = True
    for appid, name, rows, uniq, in_db, ins in report:
        match = (uniq == in_db)
        ok &= match
        print("%-10d %-32s 고유 %8d  DB %8d  %s" % (appid, name[:30], uniq, in_db, "일치" if match else "불일치!"))
    print("총 소요 %.0f초 | 전체 %s" % (time.time() - t0, "일치" if ok else "불일치 있음"))


if __name__ == "__main__":
    main()
